from docx import Document
from docx.shared import Inches
from docx.shared import Pt
from datetime import datetime
import os

# -----------------
# PLC元件資料庫
# -----------------

components = [

    {
        "name": "BW63SAG",
        "type": "無熔絲斷路器(NFB)",
        "description": """
BW63SAG主要用於保護電路，
當過電流或短路發生時會自動跳脫。
""",
        "image": "images/BW63SAG.jpg"
    },

    {
        "name": "SC-N3",
        "type": "電磁接觸器(MC)",
        "description": """
SC-N3常用於馬達控制。
PLC輸出訊號後驅動接觸器線圈，
進而控制大電流負載。
""",
        "image": "images/SC-N3.jpg"
    },

    {
        "name": "G6B-47BND",
        "type": "介面繼電器",
        "description": """
G6B系列常作為PLC與外部設備之間的隔離介面。
""",
        "image": "images/G6B-47BND.jpg"
    },

    {
        "name": "DF102",
        "type": "保險絲座",
        "description": """
搭配10x38mm保險絲使用，
提供短路保護功能。
""",
        "image": "images/DF102.jpg"
    }
]

# -----------------
# PLC I/O
# -----------------

plc_io = [

    ["X0","啟動按鈕"],
    ["X1","停止按鈕"],
    ["X2","緊急停止"],

    ["Y0","主馬達"],
    ["Y1","警示燈"],
    ["Y2","蜂鳴器"]
]

# -----------------
# 建立文件
# -----------------

doc = Document()

# -----------------
# 封面
# -----------------

title = doc.add_heading(
    "PLC控制系統學習教材",
    level=0
)

p = doc.add_paragraph()
p.add_run("作者：新人工程師\n")
p.add_run(
    datetime.now().strftime("%Y-%m-%d")
)

doc.add_page_break()

# -----------------
# 目錄提示
# -----------------

doc.add_heading(
    "目錄",
    level=1
)

doc.add_paragraph(
    "開啟Word後可參考資料 > 目錄 > 自動目錄"
)

doc.add_page_break()

# -----------------
# 第一章
# -----------------

doc.add_heading(
    "第一章 控制盤元件介紹",
    level=1
)

doc.add_paragraph(
"""
PLC控制盤由保護元件、
控制元件與執行元件所組成。
"""
)

# -----------------
# 元件章節
# -----------------

for item in components:

    doc.add_heading(
        item["name"],
        level=2
    )

    table = doc.add_table(
        rows=2,
        cols=2
    )

    table.style = "Table Grid"

    table.cell(0,0).text = "元件名稱"
    table.cell(0,1).text = item["name"]

    table.cell(1,0).text = "元件類型"
    table.cell(1,1).text = item["type"]

    doc.add_paragraph(
        item["description"]
    )

    image_path = item["image"]

    if os.path.exists(image_path):

        try:

            doc.add_picture(
                image_path,
                width=Inches(2.5)
            )

        except:
            pass

    doc.add_paragraph()

# -----------------
# 換頁
# -----------------

doc.add_page_break()

# -----------------
# 第二章
# -----------------

doc.add_heading(
    "第二章 PLC輸入輸出",
    level=1
)

doc.add_paragraph(
"""
PLC透過X接收輸入訊號，
透過Y輸出控制設備。
"""
)

# -----------------
# I/O表
# -----------------

table = doc.add_table(
    rows=1,
    cols=2
)

table.style = "Table Grid"

header = table.rows[0].cells

header[0].text = "點位"
header[1].text = "功能"

for io in plc_io:

    row = table.add_row().cells

    row[0].text = io[0]
    row[1].text = io[1]

# -----------------
# 第三章
# -----------------

doc.add_page_break()

doc.add_heading(
    "第三章 PLC控制流程",
    level=1
)

doc.add_paragraph(
"""
當按下啟動按鈕(X0)時，
PLC判斷條件成立，
輸出Y0驅動馬達運轉。

當按下停止按鈕(X1)時，
PLC關閉Y0輸出。
"""
)

# -----------------
# 學習心得
# -----------------

doc.add_page_break()

doc.add_heading(
    "第四章 學習心得",
    level=1
)

doc.add_paragraph(
"""
透過學習NFB、
電磁接觸器、
繼電器與PLC I/O，
可逐步建立工業自動化的基礎能力。
"""
)

# -----------------
# 儲存
# -----------------

filename = "PLC控制系統學習教材.docx"

doc.save(filename)

print(f"完成：{filename}")